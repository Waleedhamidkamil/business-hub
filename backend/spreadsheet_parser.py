"""
Phase M.3 - parse a clipboard-pasted spreadsheet file into a 2D string matrix
that the bulk-paste wizard can ingest.

Supports:
  • .xlsx / .xlsm - via openpyxl (already a deps for the Data Manager exports)
  • .csv          - comma-separated, dialect-sniffed
  • .tsv / .txt   - tab-separated (fallback to whatever sniff finds)

Returns a list of rows; each row is a list of strings. Trailing all-empty
rows are dropped so the wizard doesn't import a million blank lines that
Excel padded the file with.
"""
from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path


_MAX_ROWS = 5000     # safety cap - the wizard previews 8 and we never
                     # expect a Material list bigger than this.
_MAX_COLS = 64       # discard pathologically wide sheets early.


def _trim_trailing_empty(rows: list[list[str]]) -> list[list[str]]:
    while rows and not any((c or "").strip() for c in rows[-1]):
        rows.pop()
    return rows


def _looks_like_xlsx(path: Path) -> bool:
    """xlsx files are zip archives starting with the PK signature."""
    try:
        with open(path, "rb") as f:
            sig = f.read(4)
    except OSError:
        return False
    return sig[:2] == b"PK"


def _parse_xlsx(path: Path) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except Exception as e:
        raise RuntimeError(
            "openpyxl is required to parse .xlsx files. "
            "Run `pip install openpyxl` in the app's environment."
        ) from e
    if not zipfile.is_zipfile(str(path)):
        raise ValueError("File doesn't look like a valid .xlsx (not a zip archive).")

    # data_only=True so formula cells return their cached value, not the formula text.
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows: list[list[str]] = []
        # iter_rows(values_only=True) yields tuples of cell values.
        for r_i, row in enumerate(ws.iter_rows(values_only=True)):
            if r_i >= _MAX_ROWS:
                break
            cells = [_cell_to_str(v) for v in row[:_MAX_COLS]]
            rows.append(cells)
        # Pad to widest row so the grid is rectangular.
        if rows:
            w = max(len(r) for r in rows)
            for r in rows:
                while len(r) < w:
                    r.append("")
        return _trim_trailing_empty(rows)
    finally:
        try: wb.close()
        except Exception: pass


def _cell_to_str(v) -> str:
    """openpyxl returns ints, floats, datetimes, bools, None - coerce to str."""
    if v is None:
        return ""
    # Floats that happen to be whole numbers should print as ints.
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    # Default representation works for str / int / float / datetime
    return str(v)


def _parse_delimited(path: Path) -> list[list[str]]:
    """CSV / TSV / arbitrary delimited text. Sniffs the dialect from a sample."""
    # Open with utf-8-sig so a BOM doesn't pollute the first cell.
    raw = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Couldn't decode file as text - is it really a .csv/.tsv?")

    # Try to sniff the dialect; fall back to comma.
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        class _D(csv.excel):
            delimiter = "\t" if "\t" in sample else ","
        dialect = _D
    reader = csv.reader(io.StringIO(text), dialect)
    rows: list[list[str]] = []
    for r_i, row in enumerate(reader):
        if r_i >= _MAX_ROWS:
            break
        rows.append([(c or "").strip() for c in row[:_MAX_COLS]])
    if rows:
        w = max(len(r) for r in rows)
        for r in rows:
            while len(r) < w:
                r.append("")
    return _trim_trailing_empty(rows)


def parse_file(path: str | Path) -> dict:
    """
    Returns a dict the frontend can consume:
        {"ok": True,  "rows": [[...], ...], "kind": "xlsx" | "delimited"}
        {"ok": False, "error": "..."}
    Never raises - wraps everything for the JS bridge.
    """
    p = Path(str(path))
    if not p.exists() or not p.is_file():
        return {"ok": False, "error": f"File not found: {p}"}
    suffix = p.suffix.lower()
    try:
        if suffix in (".xlsx", ".xlsm") or _looks_like_xlsx(p):
            return {"ok": True, "rows": _parse_xlsx(p), "kind": "xlsx"}
        return {"ok": True, "rows": _parse_delimited(p), "kind": "delimited"}
    except Exception as e:
        return {"ok": False, "error": str(e)}
